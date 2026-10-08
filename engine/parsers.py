"""Deterministic CSV and XLSX statement parsers (no model involved)."""
from __future__ import annotations

import csv
import io
import re
from datetime import date
from pathlib import Path

from openpyxl import load_workbook

from .models import CURRENCY, Position, SourceLocation, Statement
from .numbers import EN_US, PT_BR, parse_decimal

_DATE_BR = re.compile(r"(\d{2})/(\d{2})/(\d{4})")
_DATE_ISO = re.compile(r"(\d{4})-(\d{2})-(\d{2})")


class StatementParseError(ValueError):
    pass


def parse_alfa_csv(path: Path) -> Statement:
    """Semicolon CSV, decimal comma, thousands dot, dd/mm/yyyy dates."""
    name = path.name
    rows = list(csv.reader(io.StringIO(path.read_text(encoding="utf-8-sig")), delimiter=";"))
    if not rows or not rows[0]:
        raise StatementParseError(f"{name}: empty file")
    broker = re.sub(r"\s*\(.*", "", rows[0][0].split(" - ")[0]).strip()
    account = as_of = None
    header_idx = None
    for i, row in enumerate(rows):
        first = row[0].strip() if row else ""
        if first.startswith("Cliente:"):
            account = first.split(":", 1)[1].strip()
        elif first.startswith("Data da posição:"):
            m = _DATE_BR.search(first)
            if m:
                as_of = date(int(m[3]), int(m[2]), int(m[1]))
        elif first == "Ativo":
            header_idx = i
            break
    if header_idx is None or account is None or as_of is None:
        raise StatementParseError(f"{name}: header, account or date not found")

    positions, declared_total = [], None
    for i in range(header_idx + 1, len(rows)):
        row = rows[i]
        if not row or not any(c.strip() for c in row):
            continue
        label = row[0].strip()
        if label.lower().startswith("total"):
            declared_total = parse_decimal(row[3], PT_BR)
            continue
        if len(row) < 4:
            raise StatementParseError(f"{name}: row {i + 1} has {len(row)} columns, expected 4")
        positions.append(Position(
            broker=broker, ticker=label,
            quantity=parse_decimal(row[1], PT_BR),
            avg_price=parse_decimal(row[2], PT_BR),
            declared_value=parse_decimal(row[3], PT_BR),
            source=SourceLocation(file=name, row=i + 1, excerpt=";".join(row)),
        ))
    return Statement(broker, account, as_of, CURRENCY, declared_total, positions, name)


def parse_beta_xlsx(path: Path) -> Statement:
    """English headers, ISO dates, first sheet. Read-only; formulas are never evaluated."""
    name = path.name
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb.worksheets[0]
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        sheet = ws.title
    finally:
        wb.close()
    if not rows:
        raise StatementParseError(f"{name}: empty workbook")
    broker = re.sub(r"\s*\(.*", "", str(rows[0][0]).split(" - ")[0]).strip()
    account = as_of = None
    header_idx = None
    for i, row in enumerate(rows):
        first = str(row[0]).strip() if row and row[0] is not None else ""
        if first.startswith("Account:"):
            account = first.split(":", 1)[1].strip()
        elif first.startswith("As of:"):
            m = _DATE_ISO.search(first)
            if m:
                as_of = date(int(m[1]), int(m[2]), int(m[3]))
        elif first == "Ticker":
            header_idx = i
            break
    if header_idx is None or account is None or as_of is None:
        raise StatementParseError(f"{name}: header, account or date not found")

    positions, declared_total = [], None
    for i in range(header_idx + 1, len(rows)):
        row = rows[i]
        if not row or row[0] is None:
            continue
        label = str(row[0]).strip()
        if label.upper() == "TOTAL":
            declared_total = parse_decimal(row[3], EN_US)
            continue
        positions.append(Position(
            broker=broker, ticker=label,
            quantity=parse_decimal(row[1], EN_US),
            avg_price=parse_decimal(row[2], EN_US),
            declared_value=parse_decimal(row[3], EN_US),
            source=SourceLocation(file=name, sheet=sheet, row=i + 1, column="A:D"),
        ))
    return Statement(broker, account, as_of, CURRENCY, declared_total, positions, name)
